"""Ekspor transaksi ke CSV, Excel (xlsx), dan PDF.

Berisi data TIMBANGAN saja — tidak ada kolom harga atau nilai uang."""
from __future__ import annotations

import csv
import io
from datetime import date

from app.timbangan import format as fmt
from app.timbangan import logic
from app.timbangan.service import BarisTransaksi

JUDUL = [
    "Transaction ID", "Tanggal", "Supplier", "Kendaraan", "No. Nota",
    "Timbang Masuk (kg)", "Timbang Kosong (kg)", "Netto Timbangan (kg)",
    "Jumlah Karung", "Total Tara (kg)", "Netto Kontrol (kg)", "Berat Nota (kg)",
    "Selisih (kg)", "Selisih (%)",
    "Status Kontrol", "Status Transaksi", "Diinput Oleh",
]


def _baris(b: BarisTransaksi, nama_akun: dict[int, str]) -> list:
    t = b.trx
    return [
        t.kode, str(t.tanggal), b.supplier_nama, b.nomor_polisi or "",
        t.nomor_nota or "",
        float(t.timbang_masuk_kg), float(t.timbang_kosong_kg),
        float(t.netto_timbangan_kg), int(t.jumlah_karung), float(t.total_tara_kg),
        float(t.netto_kontrol_kg), float(t.berat_nota_kg), float(t.selisih_kg),
        round(float(t.selisih_persen or 0), 2),
        logic.STATUS_LABEL.get(t.status_kontrol, t.status_kontrol),
        t.status_transaksi,
        nama_akun.get(t.dibuat_oleh_id, ""),
    ]


def ke_csv(baris: list[BarisTransaksi], nama_akun: dict[int, str]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(JUDUL)
    for b in baris:
        w.writerow(_baris(b, nama_akun))
    return buf.getvalue().encode("utf-8-sig")


def ke_xlsx(baris: list[BarisTransaksi], nama_akun: dict[int, str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Transaksi Timbang"
    ws.append(JUDUL)
    kepala = ws[1]
    for sel in kepala:
        sel.font = Font(bold=True, color="FFFFFF")
        sel.fill = PatternFill("solid", fgColor="1F3A5F")
        sel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    warna = {
        logic.STATUS_NORMAL: "E7F6EC",
        logic.STATUS_PERIKSA: "FFF1E0",
        logic.STATUS_INVESTIGASI: "FDE7E7",
    }
    for b in baris:
        ws.append(_baris(b, nama_akun))
        isi = warna.get(b.trx.status_kontrol)
        if isi:
            for sel in ws[ws.max_row]:
                sel.fill = PatternFill("solid", fgColor=isi)

    for i, judul in enumerate(JUDUL, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(24, len(judul) + 2))
    ws.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def ke_pdf(
    baris: list[BarisTransaksi],
    nama_akun: dict[int, str],
    *,
    judul: str = "Laporan Transaksi Timbang Gabah",
    subjudul: str = "",
) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=10 * mm, rightMargin=10 * mm,
        topMargin=12 * mm, bottomMargin=12 * mm,
        title=judul,
    )
    gaya = getSampleStyleSheet()
    isi = [Paragraph(judul, gaya["Title"])]
    if subjudul:
        isi.append(Paragraph(subjudul, gaya["Normal"]))
    isi.append(Spacer(1, 6 * mm))

    kepala = ["ID", "Tanggal", "Supplier", "Kendaraan", "Masuk", "Kosong",
              "Karung", "Tara", "Netto Kontrol", "Nota", "Selisih", "%", "Status"]
    data = [kepala]
    for b in baris:
        t = b.trx
        data.append([
            t.kode.replace("TRX-", ""),
            f"{t.tanggal:%d/%m/%y}",
            (b.supplier_nama or "")[:22],
            (b.nomor_polisi or "-")[:12],
            fmt.angka(t.timbang_masuk_kg),
            fmt.angka(t.timbang_kosong_kg),
            fmt.angka(t.jumlah_karung),
            fmt.angka(t.total_tara_kg),
            fmt.angka(t.netto_kontrol_kg),
            fmt.angka(t.berat_nota_kg),
            fmt.kg_tanda(t.selisih_kg),
            fmt.persen(t.selisih_persen),
            logic.STATUS_LABEL.get(t.status_kontrol, t.status_kontrol),
        ])

    tabel = Table(data, repeatRows=1)
    gaya_tabel = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("ALIGN", (4, 1), (-2, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B9C4D0")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    warna = {
        logic.STATUS_NORMAL: colors.HexColor("#E7F6EC"),
        logic.STATUS_PERIKSA: colors.HexColor("#FFF1E0"),
        logic.STATUS_INVESTIGASI: colors.HexColor("#FDE7E7"),
    }
    for i, b in enumerate(baris, start=1):
        isi_warna = warna.get(b.trx.status_kontrol)
        if isi_warna:
            gaya_tabel.append(("BACKGROUND", (0, i), (-1, i), isi_warna))
    tabel.setStyle(TableStyle(gaya_tabel))
    isi.append(tabel)

    total_netto = sum(float(b.trx.netto_kontrol_kg) for b in baris)
    total_selisih = sum(float(b.trx.selisih_kg) for b in baris)
    isi.append(Spacer(1, 5 * mm))
    isi.append(Paragraph(
        f"<b>{len(baris)} transaksi</b> &nbsp;|&nbsp; total netto kontrol "
        f"{fmt.kg(total_netto)} &nbsp;|&nbsp; total selisih "
        f"{fmt.kg_tanda(total_selisih)}",
        gaya["Normal"],
    ))
    isi.append(Spacer(1, 3 * mm))
    isi.append(Paragraph(
        "<i>Status pada laporan ini adalah alat kontrol internal, bukan "
        "kesimpulan atas perilaku supplier.</i>", gaya["Italic"],
    ))

    doc.build(isi)
    return buf.getvalue()


def nama_berkas(ekstensi: str, dari: date | None, sampai: date | None) -> str:
    bagian = ["transaksi-timbang"]
    if dari:
        bagian.append(f"{dari:%Y%m%d}")
    if sampai:
        bagian.append(f"sd-{sampai:%Y%m%d}")
    return "-".join(bagian) + "." + ekstensi
