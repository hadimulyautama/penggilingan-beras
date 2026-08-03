"""Bangun file .docx surat resmi PT HMU dengan KOP PERSIS seperti template
perusahaan (Template-Surat-Tugas.docx dkk) — kop teks, tanpa logo.

Dipakai tool MCP 'buat dokumen' agar file yang dihasilkan berkop identik dengan
dokumen resmi yang sudah ada. python-docx. MURNI (tanpa IO Drive/DB).
"""
from __future__ import annotations

import io
from datetime import date

# Kop PERSIS dari body Template-*.docx (diverifikasi 2026-08-03). Jangan diubah
# tanpa memeriksa template asli — harus identik dengan surat resmi yang beredar.
_KOP = [
    ("PT HADI MULYA UTAMA", True, 14),
    ("Tegalrejo RT/RW 21/07, Masaran, Sragen, Jawa Tengah", False, 11),
    ("Telp. (+62) 851-7416-7950   •   Email: info@hadimulya.com", False, 11),
]

_BULAN = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
          "Agustus", "September", "Oktober", "November", "Desember"]


def _tanggal_id(d: date) -> str:
    return f"Sragen, {d.day} {_BULAN[d.month]} {d.year}"


def surat_docx(
    *,
    judul: str,
    isi: str,
    nomor: str = "",
    tanggal: date | None = None,
    jabatan_ttd: str = "Direktur Utama",
    nama_ttd: str = "Ahmad Wahid Nurhani",
) -> bytes:
    """Rangkai .docx surat resmi (kop + judul + nomor + isi + tanggal + ttd).

    `isi` = badan surat (multi-baris; disusun pemanggil sesuai jenis dokumen).
    Kembalikan bytes .docx siap diunggah ke Drive.
    """
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
    from docx.shared import Pt

    doc = docx.Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # ── Kop (tengah) ──
    for teks, bold, size in _KOP:
        p = doc.add_paragraph()
        p.alignment = AL.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(teks)
        r.bold = bold
        r.font.size = Pt(size)
    # garis pemisah
    gp = doc.add_paragraph()
    gp.paragraph_format.space_before = Pt(2)
    gr = gp.add_run("_" * 78)
    gr.bold = True

    # ── Judul + nomor (tengah) ──
    pj = doc.add_paragraph()
    pj.alignment = AL.CENTER
    pj.paragraph_format.space_before = Pt(6)
    rj = pj.add_run(judul.upper())
    rj.bold = True
    rj.underline = True
    rj.font.size = Pt(12)
    if nomor:
        pn = doc.add_paragraph()
        pn.alignment = AL.CENTER
        pn.paragraph_format.space_after = Pt(10)
        pn.add_run(f"Nomor : {nomor}")

    # ── Isi (kiri, per baris) ──
    for baris in isi.replace("\r\n", "\n").split("\n"):
        p = doc.add_paragraph(baris)
        p.paragraph_format.space_after = Pt(2)

    # ── Tanggal + blok tanda tangan (kanan) ──
    d = tanggal or date.today()
    for teks in ("", _tanggal_id(d), "PT Hadi Mulya Utama", f"{jabatan_ttd},", "", "", ""):
        p = doc.add_paragraph()
        p.alignment = AL.RIGHT
        p.paragraph_format.space_after = Pt(0)
        if teks:
            p.add_run(teks)
    pt = doc.add_paragraph()
    pt.alignment = AL.RIGHT
    pt.add_run(nama_ttd).bold = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
